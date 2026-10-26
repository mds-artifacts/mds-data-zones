"""
Step 1 - Read every extraction tab of the SLR workbook into ONE tidy table.

The extraction tabs were filled at different times, so their headers differ in
spelling ("Ground thruth From", "trianing envrionment data temporality", ...)
and the order of columns is not the same. This module maps each tab onto one
canonical schema by matching header text, with explicit overrides for tabs
whose cells do not sit under the header they belong to.

Output: data/studies_raw.csv  (one row per extracted primary study, per tab)
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data" / "Final_SLR_extractions.xlsx"

# Canonical schema -> regexes tried against the normalised header, in order.
FIELDS: dict[str, list[str]] = {
    "ps_id": [r"^ps ?#", r"^ref ?#"],
    "title": [r"^title$", r"^detection model used"],
    "author": [r"^author"],
    "year": [r"^year"],
    "venue": [r"^venue"],
    "approach": [r"detection approach"],
    "model_class": [r"detection model class"],
    "model": [r"detection model\(primary"],
    "learning": [r"learning technique"],
    "architecture": [r"model architecture"],
    "other_techniques": [r"other techniques"],
    "baselines": [r"baseline codes"],
    "knowledge_zone": [r"knowledge zone"],
    "ground_truth": [r"ground (truth|thruth|thurth) (labels )?from", r"^ground (truth|thruth)$"],
    "datasets": [r"datasets name", r"claims set"],
    "train_temporality": [r"^trai\w* .*temp", r"^train env", r"^trained e", r"trained set"],
    "test_temporality": [r"^test\w* .*temp", r"^test env", r"^tested env", r"^test set", r"^test temp"],
    "dataset_temporality": [r"^dataset temporality"],
    "domains": [r"^domains"],
    "factcheck_orgs": [r"professional fact"],
    "reliable_sources": [r"mainstream media"],
    "features": [r"feature extraction"],
    "reasoning": [r"reasoning"],
    "findings": [r"^findings"],
    "metrics": [r"metrics codes"],
    "labels": [r"^lables|^labels"],
    "evaluation": [r"^evaluation"],
}
SWOT = ["S", "W", "O", "T"]

# Tabs whose cells are shifted relative to their headers (verified by hand):
# the values follow the coding order used when the tab was generated, so we
# map them by position instead of by header.
POSITIONAL = {
    "7 Extraction - Mohommad": {
        "ps_id": 0, "title": 1, "author": 2, "year": 3, "venue": 4,
        "approach": 5, "model_class": 6, "model": 7, "learning": 8,
        "architecture": 9, "other_techniques": 10, "baselines": 11,
        "knowledge_zone": 12, "ground_truth": 13, "datasets": 14,
        "dataset_temporality": 15, "domains": 16, "factcheck_orgs": 17,
        "reliable_sources": 18, "features": 19, "reasoning": 20,
        "findings": 21, "metrics": 22, "labels": 23, "evaluation": 24,
        "S": 25, "W": 26, "O": 27, "T": 28,
    },
}

# Short label for each source review (used in figures / tables).
REVIEW_LABEL = {
    "5": "Bodaghi", "6": "Tabrani", "7": "Mohommad", "8": "Fanlau",
    "9": "Bashaddadh", "10": "Darnoto", "11": "Elias", "12": "Alkomah & Sheldon",
    "13": "Gueddoudj", "14": "Xu", "15": "Tajrian", "16": "LLM survey",
    "17": "Thomas",
}


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "")).strip().lower()


def _clean(v):
    if v is None:
        return None
    s = str(v).replace(" ", " ").strip()
    return None if s in ("", "None", "-", "—", "N/A", "n/a", "NA") else s


def extraction_sheets(wb) -> list[str]:
    out = []
    for name in wb.sheetnames:
        n = name.lower()
        if "extraction" in n and " ie " not in f" {n} " and not n.startswith("ie") and "- ie" not in n:
            out.append(name)
    return out


def _resolve_columns(headers: list[str], sheet: str) -> dict[str, int]:
    if sheet in POSITIONAL:
        return dict(POSITIONAL[sheet])
    cols: dict[str, int] = {}
    used: set[int] = set()
    for field, pats in FIELDS.items():
        for p in pats:
            hit = next((i for i, h in enumerate(headers) if i not in used and re.search(p, h)), None)
            if hit is not None:
                cols[field] = hit
                used.add(hit)
                break
    # SWOT: either four headed columns S/W/O/T, or a "MY ANALYSIS and SWOT"
    # column followed by three unnamed columns. Prefer the coded block.
    swot_block = next((i for i, h in enumerate(headers) if "swot" in h), None)
    if swot_block is not None:
        for k, f in enumerate(SWOT):
            cols[f] = swot_block + k
    else:
        for f in SWOT:
            hit = next((i for i, h in enumerate(headers) if h == f.lower()), None)
            if hit is not None:
                cols[f] = hit
    return cols


def load(workbook: Path = WORKBOOK) -> pd.DataFrame:
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    frames = []
    for sheet in extraction_sheets(wb):
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        headers = [_norm(h) for h in rows[0]]
        cols = _resolve_columns(headers, sheet)
        if "ps_id" not in cols:   # e.g. Bodaghi: PS column carries a wrong header - detect it by its values
            body = [r for r in rows[1:] if r and sum(c not in (None, "") for c in r) >= 4]
            for i in range(min(3, len(headers))):
                vals = [str(r[i]) for r in body if i < len(r) and r[i] is not None]
                if vals and sum(bool(re.fullmatch(r"\s*PS\s*\d+\s*", v)) for v in vals) >= 0.8 * len(vals):
                    cols["ps_id"] = i
                    break
        recs = []
        for r in rows[1:]:
            if not r or sum(c not in (None, "") for c in r) < 4:
                continue  # blank / note rows
            rec = {f: _clean(r[i]) if i < len(r) else None for f, i in cols.items()}
            recs.append(rec)
        df = pd.DataFrame(recs)
        num = re.match(r"\s*(\d+)", sheet)
        df.insert(0, "review_no", num.group(1) if num else sheet)
        df.insert(1, "review", REVIEW_LABEL.get(num.group(1) if num else "", sheet))
        df.insert(2, "sheet", sheet)
        frames.append(df)
    out = pd.concat(frames, ignore_index=True)
    # Single-label temporality tabs: the one label describes the whole design.
    single = out["dataset_temporality"].notna() if "dataset_temporality" in out else False
    out["temporality_single_label"] = single
    for c in ("train_temporality", "test_temporality"):
        out[c] = out[c].where(out[c].notna(), out.get("dataset_temporality"))
    return out


def load_primary(workbook: Path = WORKBOOK) -> pd.DataFrame:
    """load() minus the rows dropped in config.py (duplicates / non-primary), so every figure uses the same N."""
    import config as C
    df = load(workbook)
    seq = df.groupby("review").cumcount() + 1
    uid = df["review"] + "#" + seq.astype(str)
    drop = set(getattr(C, "DROP_DUPLICATES", [])) | set(getattr(C, "EXCLUDE_NOT_PRIMARY", []))
    df = df[~uid.isin(drop)].reset_index(drop=True)
    # tidy PS# typos such as "P136" -> "PS136"
    df["ps_id"] = df["ps_id"].map(lambda v: re.sub(r"^\s*PS?\s*(\d+)\s*$", r"PS\1", str(v)) if v is not None else v)
    return df


if __name__ == "__main__":
    df = load()
    (ROOT / "data").mkdir(exist_ok=True)
    df.to_csv(ROOT / "data" / "studies_raw.csv", index=False)
    print(df.groupby("review", sort=False).size().to_string())
    print("TOTAL rows:", len(df))
    print("Missing by field (%):")
    print((df.isna().mean() * 100).round(0).sort_values(ascending=False).head(15).to_string())
